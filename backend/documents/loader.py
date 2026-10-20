"""Normalize supported files into citation-friendly pages or worksheet sections."""

from pathlib import Path
from zipfile import ZipFile

from .pdf_loader import PageText, PDFDocument, load_pdf

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls", ".csv", ".txt", ".md"}
MAX_TEXT_CHARS = 5_000_000


def load_document(file_path: str | Path) -> PDFDocument:
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("Supported formats: PDF, DOCX, XLSX, XLS, CSV, TXT, MD.")
    if ext == ".pdf":
        result = load_pdf(str(path))
        if not any(p.text.strip() for p in result.pages):
            raise ValueError("No readable text found. Scanned PDFs need OCR before upload.")
        return result
    sections: list[str] = []
    if ext in {".docx", ".xlsx"}:
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100_000_000:
                raise ValueError("Expanded document exceeds the 100 MB processing limit.")
    if ext == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(str(path))
        parts = []
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                parts.append(block.text)
            elif isinstance(block, Table):
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        text = "\n".join(parts)
        sections = [text[i : i + 12000] for i in range(0, len(text), 12000)]
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book:
                if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > 500_000:
                    raise ValueError("Worksheet exceeds the 500,000 cell limit.")
                lines = []
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if any(v is not None for v in row):
                        lines.append(f"Row {number}: " + " | ".join("" if v is None else str(v) for v in row))
                for start in range(0, len(lines), 150):
                    sections.append(f"Worksheet: {sheet.title}\n" + "\n".join(lines[start : start + 150]))
        finally:
            book.close()
    elif ext == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows * sheet.ncols > 500_000:
                    raise ValueError("Worksheet exceeds the 500,000 cell limit.")
                for start in range(0, sheet.nrows, 150):
                    lines = [
                        f"Row {i + 1}: " + " | ".join(str(v) for v in sheet.row_values(i))
                        for i in range(start, min(start + 150, sheet.nrows))
                    ]
                    sections.append(f"Worksheet: {sheet.name}\n" + "\n".join(lines))
        finally:
            book.release_resources()
    else:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        if ext == ".csv":
            import csv
            import io

            text = "\n".join(f"Row {i}: " + " | ".join(row) for i, row in enumerate(csv.reader(io.StringIO(text)), 1))
        sections = [text[i : i + 12000] for i in range(0, len(text), 12000)]
    if sum(map(len, sections)) > MAX_TEXT_CHARS:
        raise ValueError("Document exceeds the 5 million character limit.")
    if not any(s.strip() for s in sections):
        raise ValueError("No readable content found in this file.")
    return PDFDocument(
        name=path.name,
        file_path=str(path),
        pages=[
            PageText(page_number=i, text=text, doc_name=path.name, total_pages=len(sections))
            for i, text in enumerate(sections, 1)
        ],
    )
