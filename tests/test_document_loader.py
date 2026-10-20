from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from backend.documents.loader import load_document


def test_word_preserves_paragraph_and_table_order(tmp_path):
    path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("Before the table")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "5000 USD"
    doc.add_paragraph("After the table")
    doc.save(path)
    loaded = load_document(path)
    text = loaded.pages[0].text
    assert text.index("Before") < text.index("Revenue | 5000 USD") < text.index("After")
    assert loaded.get_page(1).doc_name == "report.docx"


def test_excel_preserves_sheets_rows_formulas_and_zero(tmp_path):
    path = tmp_path / "budget.xlsx"
    book = Workbook()
    first = book.active
    first.title = "Revenue"
    first.append(["Quarter", "Amount"])
    first.append(["Q1", 0])
    first.append(["Q2", "=SUM(B2:B2)"])
    second = book.create_sheet("Costs")
    second.append(["Hosting", 900])
    book.save(path)
    loaded = load_document(path)
    assert loaded.total_pages == 2
    assert "Worksheet: Revenue" in loaded.pages[0].text
    assert "Row 2: Q1 | 0" in loaded.pages[0].text
    assert "=SUM(B2:B2)" in loaded.pages[0].text
    assert "Worksheet: Costs" in loaded.pages[1].text
    assert "900" in loaded.pages[1].text
    assert loaded.get_page(2).page_number == 2


def test_large_sheet_sections_keep_original_row_numbers(tmp_path):
    path = tmp_path / "rows.xlsx"
    book = Workbook()
    for i in range(151):
        book.active.append([i])
    book.save(path)
    loaded = load_document(path)
    assert loaded.total_pages == 2
    assert "Row 151: 150" in loaded.pages[1].text


def test_csv_quoted_cells_and_unicode(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text('\ufeffName,Note,Amount\nAyoub,"Hello, world",0\n', encoding="utf-8")
    result = load_document(path)
    assert "Row 2: Ayoub | Hello, world | 0" in result.pages[0].text


@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_text_and_markdown(tmp_path, suffix):
    path = tmp_path / ("notes" + suffix)
    path.write_text("Document findings: café — 2026", encoding="utf-8")
    assert "café — 2026" in load_document(path).pages[0].text


def test_pdf_uses_existing_extractor():
    loaded = load_document(Path(__file__).parents[1] / "sample.pdf")
    assert loaded.total_pages > 0
    assert any(page.text.strip() for page in loaded.pages)


@pytest.mark.parametrize("suffix", [".txt", ".docx", ".xlsx"])
def test_empty_files_rejected(tmp_path, suffix):
    path = tmp_path / ("empty" + suffix)
    if suffix == ".docx":
        Document().save(path)
    elif suffix == ".xlsx":
        Workbook().save(path)
    else:
        path.write_text("   ")
    with pytest.raises(ValueError, match="No readable"):
        load_document(path)


def test_unsupported_extension_rejected(tmp_path):
    with pytest.raises(ValueError, match="Supported formats"):
        load_document(tmp_path / "image.png")
