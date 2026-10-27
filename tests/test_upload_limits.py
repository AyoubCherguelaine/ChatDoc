"""Upload limits apply before storage, embeddings, or OCR across API entrypoints."""

from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from backend import rag, server
from backend.auth import get_current_user
from backend.documents.limits import (
    MAX_UPLOAD_BYTES, UploadSizeError, validate_pdf_pages, validate_upload_size,
)
from backend.documents.loader import load_document
from backend.storage import storage_service
from backend.v1 import routes


def pdf_bytes(pages):
    stream = BytesIO()
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=100, height=100)
    writer.write(stream)
    return stream.getvalue()


@pytest.fixture
def client(monkeypatch, tmp_path):
    server.app.dependency_overrides[get_current_user] = lambda: {"id": "limit-test-user"}
    server.app.dependency_overrides[routes.get_company_or_api_auth] = lambda: SimpleNamespace(user_id="limit-test-user")
    monkeypatch.setattr(rag.settings, "UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setattr(rag.settings, "INDEX_DIR", str(tmp_path / "indexes"))
    monkeypatch.setattr(storage_service, "upload", MagicMock(side_effect=AssertionError("Must not store rejected files")))
    monkeypatch.setattr(rag, "SmartEmbedding", MagicMock(side_effect=AssertionError("Must not embed rejected files")))
    try:
        yield TestClient(server.app)
    finally:
        server.app.dependency_overrides.clear()


def test_size_boundary_includes_exactly_five_mb():
    validate_upload_size(MAX_UPLOAD_BYTES)
    for size in (0, MAX_UPLOAD_BYTES + 1):
        with pytest.raises(UploadSizeError):
            validate_upload_size(size)


def test_page_boundary_and_invalid_pdf():
    validate_pdf_pages(pdf_bytes(3))
    with pytest.raises(ValueError, match="maximum limit of 3 pages"):
        validate_pdf_pages(pdf_bytes(4))
    with pytest.raises(ValueError, match="could not be read"):
        validate_pdf_pages(b"invalid pdf")


@pytest.mark.parametrize("endpoint", ["/api/documents/upload", "/api/v1/documents/upload", "/api/v1/ocr/advanced"])
def test_every_file_upload_rejects_over_five_mb(client, endpoint):
    response = client.post(endpoint, files={"file": ("large.pdf", b"x" * (MAX_UPLOAD_BYTES + 1), "application/pdf")})
    assert response.status_code == 413
    assert "5 MB" in response.json()["detail"]


@pytest.mark.parametrize("endpoint", ["/api/documents/upload", "/api/v1/documents/upload", "/api/v1/ocr/advanced"])
def test_every_file_upload_rejects_four_page_pdf(client, endpoint):
    response = client.post(endpoint, files={"file": ("four.pdf", pdf_bytes(4), "application/pdf")}, data={"start_page": 1, "end_page": 1})
    assert response.status_code == 422
    assert "maximum limit of 3 pages" in response.json()["detail"]
    storage_service.upload.assert_not_called()
    rag.SmartEmbedding.assert_not_called()


@pytest.mark.parametrize("max_pages", [None, "0", "-1", "100"])
def test_client_cannot_raise_or_disable_page_limit(client, max_pages):
    response = client.post("/api/documents/upload", files={"file": ("four.pdf", pdf_bytes(4))}, data={} if max_pages is None else {"max_pages": max_pages})
    assert response.status_code == 422
    assert "maximum limit of 3 pages" in response.json()["detail"]


@pytest.mark.parametrize("json_upload", [True, False])
def test_raw_pipeline_upload_obeys_size_limit(client, json_upload):
    text = "x" * (MAX_UPLOAD_BYTES + 1)
    if json_upload:
        response = client.post("/api/v1/documents/upload", json={"content": text, "filename": "large.txt"})
    else:
        response = client.post("/api/v1/documents/upload", content=text, headers={"Content-Type": "text/plain", "X-Filename": "large.txt"})
    assert response.status_code == 413


def test_non_pdf_sections_obey_page_limit(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("x" * 36000)
    assert load_document(path).total_pages == 3
    path.write_text("x" * 36001)
    with pytest.raises(ValueError, match="maximum limit of 3 pages"):
        load_document(path)


def test_spreadsheet_sheet_count_obeys_page_limit(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "sheets.xlsx"
    book = Workbook()
    book.active.append(["First sheet"])
    for name in ("Second", "Third"):
        book.create_sheet(name).append([name])
    book.save(path)
    assert load_document(path).total_pages == 3
    book.create_sheet("Fourth").append(["Fourth"])
    book.save(path)
    with pytest.raises(ValueError, match="maximum limit of 3 pages"):
        load_document(path)


def test_json_pipeline_rejects_over_three_sections_before_embedding(client):
    response = client.post("/api/v1/documents/upload", json={"content": "x" * 36001, "filename": "notes.txt"})
    assert response.status_code == 422
    storage_service.upload.assert_not_called()
    rag.SmartEmbedding.assert_not_called()


def test_existing_ocr_document_cannot_bypass_page_limit(client, tmp_path, monkeypatch):
    path = tmp_path / "four.pdf"
    path.write_bytes(pdf_bytes(4))
    monkeypatch.setattr(routes.ephemeral_store, "get_and_touch", lambda *a, **kw: SimpleNamespace(file_path=path, filename=path.name, doc_id="existing"))
    response = client.post("/api/v1/ocr/advanced", json={"document_id": "existing", "start_page": 1, "end_page": 1})
    assert response.status_code == 422
    assert "maximum limit of 3 pages" in response.json()["detail"]
